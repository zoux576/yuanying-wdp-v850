from __future__ import annotations
from pathlib import Path
from typing import Any
from io import BytesIO
import hashlib, json, re, zipfile, xml.etree.ElementTree as ET
from collections import Counter
from PIL import Image
from docx import Document
from docx.shared import Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from .core import connect, sha256_file, get_asset, SLOT_KIND_MAP, current_video_prompt, list_assets, now_iso, invalidate_document

EMU_PER_IN = 914400
NS = {
    'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
    'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
    'wp': 'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing',
}


def text_hash(s: str) -> str:
    return hashlib.sha256(' '.join(s.split()).encode('utf-8')).hexdigest()[:20]


def docx_sha(path: Path) -> str:
    return sha256_file(path)


def infer_scene(text: str, heading: list[str]) -> str | None:
    joined = ' '.join(heading + [text])
    m = re.search(r'\bS(\d{2})\b', joined, re.I)
    if m:
        return 'S' + m.group(1)
    m = re.search(r'第\s*([一二三四五六七八九十百\d]+)\s*场', joined)
    if m and m.group(1).isdigit():
        return f"S{int(m.group(1)):02d}"
    return None


def semantic_tag(text: str) -> str:
    t = ' '.join(text.split())
    if re.search(r'(正式参考图片|正式图片|S\d{2}-IMG-\d{2})', t, re.I):
        return 'FORMAL_SLOT'
    if re.search(r'(最终视频提示词|FINAL[- ]?VIDEO[- ]?PROMPT)', t, re.I):
        return 'VIDEO_PROMPT'
    if re.search(r'(完整故事版|故事版总板|STORYBOARD[- ]?BOARD)', t, re.I):
        return 'STORYBOARD_SLOT'
    if re.search(r'(故事版镜头清单|SHOTLIST)', t, re.I):
        return 'SHOTLIST'
    if re.search(r'(概念附录|CONCEPT APPENDIX)', t, re.I):
        return 'CONCEPT_APPENDIX'
    return 'GENERIC'


def map_anchors(docx_path: Path) -> list[dict[str, Any]]:
    doc = Document(docx_path)
    source_sha = docx_sha(docx_path)
    rows = []
    heading: list[str] = []
    seen_paths = set()
    for pi, p in enumerate(doc.paragraphs):
        text = ' '.join(p.text.split())
        if not text:
            continue
        if p.style and p.style.name.startswith('Heading'):
            level = 1
            m = re.search(r'(\d+)', p.style.name)
            if m:
                level = max(1, int(m.group(1)))
            heading = heading[:level-1] + [text]
        tag = semantic_tag(text)
        if tag != 'GENERIC' or (p.style and p.style.name.startswith('Heading')):
            path = f'body/p[{pi}]'
            aid = 'A-' + hashlib.sha256(f'{source_sha}|{path}|{text_hash(text)}'.encode()).hexdigest()[:22]
            rows.append({'anchor_id': aid, 'container_type': 'PARAGRAPH', 'container_path': path,
                         'scene_id': infer_scene(text, heading), 'semantic_tag': tag, 'exact_text': text,
                         'text_hash': text_hash(text), 'heading_path': list(heading), 'source_docx_sha': source_sha})
            seen_paths.add(path)
    for ti, table in enumerate(doc.tables):
        for ri, row in enumerate(table.rows):
            for ci, cell in enumerate(row.cells):
                local_heading = list(heading)
                for pi, p in enumerate(cell.paragraphs):
                    text = ' '.join(p.text.split())
                    if not text:
                        continue
                    tag = semantic_tag(text)
                    if tag == 'GENERIC':
                        continue
                    path = f'body/tbl[{ti}]/tr[{ri}]/tc[{ci}]/p[{pi}]'
                    if path in seen_paths:
                        continue
                    aid = 'A-' + hashlib.sha256(f'{source_sha}|{path}|{text_hash(text)}'.encode()).hexdigest()[:22]
                    rows.append({'anchor_id': aid, 'container_type': 'TABLE_CELL_PARAGRAPH', 'container_path': path,
                                 'scene_id': infer_scene(text, local_heading), 'semantic_tag': tag, 'exact_text': text,
                                 'text_hash': text_hash(text), 'heading_path': local_heading, 'source_docx_sha': source_sha})
    return rows


def save_anchors(db_path: Path, project_id: str, docx_path: Path, out_json: Path) -> list[dict[str, Any]]:
    rows = map_anchors(docx_path)
    if not rows:
        raise ValueError('no usable semantic anchors found in source DOCX')
    out_json.parent.mkdir(parents=True, exist_ok=True)
    out_json.write_text(json.dumps(rows, ensure_ascii=False, indent=2), encoding='utf-8')
    with connect(db_path) as con:
        con.execute('DELETE FROM anchors WHERE project_id=?', (project_id,))
        for a in rows:
            con.execute('INSERT INTO anchors VALUES(?,?,?,?,?,?,?,?,?,?)',
                        (project_id, a['anchor_id'], a['container_type'], a['container_path'], a['scene_id'], a['semantic_tag'],
                         a['exact_text'], a['text_hash'], json.dumps(a['heading_path'], ensure_ascii=False), a['source_docx_sha']))
        con.execute('UPDATE projects SET anchor_map_path=? WHERE project_id=?', (str(out_json), project_id))
    return rows




def source_docx_preflight(docx_path: Path) -> dict[str, Any]:
    doc = Document(docx_path)
    text = '\n'.join(p.text for p in doc.paragraphs)
    markers = {
        'PROMPT-COMPILED': len(re.findall(r'PROMPT-COMPILED', text, re.I)),
        'PRECOMPILED-DRAFT': len(re.findall(r'PRECOMPILED-DRAFT', text, re.I)),
        'ASSET-PRECOMPILED': len(re.findall(r'ASSET-PRECOMPILED', text, re.I)),
        'NON_FORMAL_LABEL': len(re.findall(r'不是正式图片|不是真正正式图片|不得上传为P18参考|FORMAL IMAGE BLOCKED', text, re.I)),
    }
    blocked = sum(markers.values()) > 0
    with zipfile.ZipFile(docx_path) as z:
        media_count = len([n for n in z.namelist() if n.startswith('word/media/') and not n.endswith('/')])
    return {
        'status': 'BLOCKED' if blocked else 'PASS',
        'markers': markers,
        'media_count': media_count,
        'failed_reasons': [f'{k}={v}' for k, v in markers.items() if v],
        'instruction': 'Upload a clean original/source DOCX; do not use a PARTIAL output containing prompt cards or draft video prompts as the production base.' if blocked else 'Source DOCX accepted.',
    }

def register_source_docx_media(db_path: Path, storage_root: Path, project_id: str, docx_path: Path) -> list[dict[str, Any]]:
    """Register pre-existing source DOCX images as immutable baseline media.

    These assets are preserved and auditable but never count as formal images or storyboard boards.
    """
    out = []
    target_dir = storage_root / project_id / 'source' / 'baseline_media'
    target_dir.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(docx_path) as z, connect(db_path) as con:
        existing = {r['sha256']: r['asset_id'] for r in con.execute("SELECT asset_id,sha256 FROM assets WHERE project_id=?", (project_id,))}
        idx = 0
        for member in sorted(n for n in z.namelist() if n.startswith('word/media/') and not n.endswith('/')):
            raw = z.read(member)
            sha = hashlib.sha256(raw).hexdigest()
            if sha in existing:
                continue
            idx += 1
            ext = Path(member).suffix.lower() or '.bin'
            asset_id = f'BASE-DOC-IMG-{idx:03d}'
            dest = target_dir / f'{asset_id}{ext}'
            dest.write_bytes(raw)
            width = height = None
            try:
                with Image.open(BytesIO(raw)) as im:
                    width, height = im.size
            except Exception:
                pass
            ts = now_iso()
            con.execute("""INSERT INTO assets(project_id,asset_id,scene_id,kind,source_kind,qc_status,file_path,sha256,width,height,mime_type,original_name,orig_fig_no,shot_id,timecode,prompt_version,allowed_for_embed,allowed_for_recompile,allowed_for_board,evidence_json,flags_json,created_at,updated_at)
                         VALUES(?,?,NULL,'DOCUMENT','USER_UPLOADED','BASELINE-PRESERVED',?,?,?,?,?,?,NULL,NULL,NULL,NULL,1,0,0,?, '[]',?,?)""",
                        (project_id, asset_id, str(dest), sha, width, height, None, Path(member).name, json.dumps({'source_docx_member': member, 'baseline': True}, ensure_ascii=False), ts, ts))
            out.append({'asset_id': asset_id, 'member': member, 'sha256': sha, 'pixels': [width, height]})
    return out

def resolve_paragraph(doc: Document, path: str) -> Paragraph:
    m = re.fullmatch(r'body/p\[(\d+)\]', path)
    if m:
        return doc.paragraphs[int(m.group(1))]
    m = re.fullmatch(r'body/tbl\[(\d+)\]/tr\[(\d+)\]/tc\[(\d+)\]/p\[(\d+)\]', path)
    if m:
        ti, ri, ci, pi = map(int, m.groups())
        return doc.tables[ti].rows[ri].cells[ci].paragraphs[pi]
    raise ValueError(f'unsupported anchor path: {path}')


def insert_after(paragraph: Paragraph, text: str | None = None) -> Paragraph:
    new_p = OxmlElement('w:p')
    paragraph._p.addnext(new_p)
    p = Paragraph(new_p, paragraph._parent)
    if text is not None:
        p.add_run(text)
    return p


def insert_before(paragraph: Paragraph, text: str | None = None) -> Paragraph:
    new_p = OxmlElement('w:p')
    paragraph._p.addprevious(new_p)
    p = Paragraph(new_p, paragraph._parent)
    if text is not None:
        p.add_run(text)
    return p


def set_no_compression(doc: Document):
    settings = doc.settings._element
    if settings.find(qn('w:doNotCompressPictures')) is None:
        settings.append(OxmlElement('w:doNotCompressPictures'))
    dpi = settings.find(qn('w14:defaultImageDpi'))
    if dpi is None:
        dpi = OxmlElement('w14:defaultImageDpi')
        settings.append(dpi)
    dpi.set(qn('w14:val'), '330')


def set_section_end(paragraph: Paragraph, landscape: bool, margin_twips: int = 680):
    pPr = paragraph._p.get_or_add_pPr()
    old = pPr.find(qn('w:sectPr'))
    if old is not None:
        pPr.remove(old)
    sect = OxmlElement('w:sectPr')
    typ = OxmlElement('w:type')
    typ.set(qn('w:val'), 'nextPage')
    sect.append(typ)
    pg = OxmlElement('w:pgSz')
    if landscape:
        pg.set(qn('w:w'), '16838')
        pg.set(qn('w:h'), '11906')
        pg.set(qn('w:orient'), 'landscape')
    else:
        pg.set(qn('w:w'), '11906')
        pg.set(qn('w:h'), '16838')
    sect.append(pg)
    mar = OxmlElement('w:pgMar')
    for k, v in [('top', margin_twips), ('right', margin_twips), ('bottom', margin_twips), ('left', margin_twips), ('header', 360), ('footer', 360), ('gutter', 0)]:
        mar.set(qn('w:' + k), str(v))
    sect.append(mar)
    pPr.append(sect)


def add_page_break_after(p: Paragraph) -> Paragraph:
    q = insert_after(p)
    q.add_run().add_break(WD_BREAK.PAGE)
    return q


def _anchor_map(db_path: Path, project_id: str) -> dict[str, dict[str, Any]]:
    with connect(db_path) as con:
        return {r['anchor_id']: dict(r) for r in con.execute('SELECT * FROM anchors WHERE project_id=?', (project_id,))}


def validate_embed_plan(db_path: Path, project_id: str, operations: list[dict[str, Any]], require_complete: bool = False) -> None:
    anchors = _anchor_map(db_path, project_id)
    seen_orders = set()
    seen_assets = []
    seen_scene_slot = set()
    for op in operations:
        order = op.get('op_order')
        if order in seen_orders:
            raise ValueError('duplicate op_order')
        seen_orders.add(order)
        aid = op.get('anchor_id')
        if aid not in anchors:
            raise ValueError(f'anchor not found: {aid}')
        anchor = anchors[aid]
        if op.get('insert_relation', 'AFTER') not in {'AFTER', 'BEFORE'}:
            raise ValueError('insert_relation must be AFTER or BEFORE')
        if op.get('orientation', 'INHERIT') not in {'INHERIT', 'PORTRAIT', 'LANDSCAPE'}:
            raise ValueError('invalid orientation')
        scene_id = op.get('scene_id')
        if scene_id and anchor.get('scene_id') and scene_id != anchor['scene_id']:
            raise ValueError(f"anchor scene mismatch: {scene_id} vs {anchor['scene_id']}")
        if op['op_type'] == 'IMAGE':
            asset = get_asset(db_path, project_id, op['asset_id'])
            expected = SLOT_KIND_MAP.get(op.get('slot_kind'))
            if not expected or asset['kind'] != expected:
                raise ValueError(f"slot mismatch {op.get('slot_kind')} -> {asset['kind']}")
            if scene_id and asset.get('scene_id') and scene_id != asset['scene_id']:
                raise ValueError(f"asset scene mismatch: {asset['asset_id']}")
            if not asset['allowed_for_embed']:
                raise ValueError(f"asset not allowed_for_embed: {asset['asset_id']}")
            if not Path(asset['file_path']).exists() or sha256_file(Path(asset['file_path'])) != asset['sha256']:
                raise ValueError(f"asset file/hash invalid: {asset['asset_id']}")
            if asset['asset_id'] in seen_assets:
                raise ValueError(f"duplicate asset in embed plan: {asset['asset_id']}")
            seen_assets.append(asset['asset_id'])
            key = (scene_id, op.get('slot_kind'), asset['asset_id'])
            if key in seen_scene_slot:
                raise ValueError(f"duplicate scene slot operation: {key}")
            seen_scene_slot.add(key)
        elif op['op_type'] == 'VIDEO_PROMPT_TEXT':
            if not scene_id:
                raise ValueError('VIDEO_PROMPT_TEXT requires scene_id')
            current_video_prompt(db_path, project_id, scene_id)
        else:
            raise ValueError(f"unsupported op_type {op['op_type']}")
    if require_complete:
        required_assets = []
        with connect(db_path) as con:
            scenes = [r['scene_id'] for r in con.execute('SELECT scene_id FROM scenes WHERE project_id=?', (project_id,))]
        for sid in scenes:
            formals = [a['asset_id'] for a in list_assets(db_path, project_id, sid, 'FORMAL_IMAGE') if a['allowed_for_embed']]
            boards = [a['asset_id'] for a in list_assets(db_path, project_id, sid, 'STORYBOARD_BOARD') if a['allowed_for_embed']]
            required_assets.extend(formals)
            if len(boards) != 1:
                raise ValueError(f'{sid}: exactly one embed-ready storyboard board required, found {len(boards)}')
            required_assets.extend(boards)
            if not any(op['op_type'] == 'VIDEO_PROMPT_TEXT' and op.get('scene_id') == sid for op in operations):
                raise ValueError(f'{sid}: FINAL video prompt operation missing')
        if Counter(seen_assets) != Counter(required_assets):
            missing = sorted((Counter(required_assets) - Counter(seen_assets)).elements())
            extra = sorted((Counter(seen_assets) - Counter(required_assets)).elements())
            raise ValueError(f'embed plan asset set mismatch; missing={missing}, extra={extra}')


def store_embed_plan(db_path: Path, project_id: str, operations: list[dict[str, Any]]) -> None:
    validate_embed_plan(db_path, project_id, operations, require_complete=False)
    with connect(db_path) as con:
        con.execute('DELETE FROM embed_plan WHERE project_id=?', (project_id,))
        for op in sorted(operations, key=lambda x: x['op_order']):
            con.execute('INSERT INTO embed_plan VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)', (
                project_id, op['op_order'], op['anchor_id'], op['op_type'], op.get('asset_id'), op.get('scene_id'),
                op.get('slot_kind'), op.get('insert_relation', 'AFTER'), op.get('orientation', 'INHERIT'),
                op.get('width_cm'), op.get('caption'), op.get('title'), int(op.get('page_break_before', False)),
                int(op.get('page_break_after', False)), op.get('text_content')
            ))
    invalidate_document(db_path, project_id, 'embed plan updated')


def build_docx(db_path: Path, project_id: str, source_docx: Path, out: Path) -> Path:
    with connect(db_path) as con:
        p = con.execute('SELECT * FROM projects WHERE project_id=?', (project_id,)).fetchone()
        ops = [dict(r) for r in con.execute('SELECT * FROM embed_plan WHERE project_id=? ORDER BY op_order', (project_id,))]
        anchors = {r['anchor_id']: dict(r) for r in con.execute('SELECT * FROM anchors WHERE project_id=?', (project_id,))}
    if not p or p['source_docx_sha'] != docx_sha(source_docx):
        raise ValueError('source DOCX hash changed; remap anchors')
    validate_embed_plan(db_path, project_id, ops, require_complete=True)
    doc = Document(source_docx)
    set_no_compression(doc)
    cursors = {}
    anchor_objects = {aid: resolve_paragraph(doc, a['container_path']) for aid, a in anchors.items()}
    for op in ops:
        a = anchors[op['anchor_id']]
        anchor = anchor_objects[op['anchor_id']]
        if text_hash(anchor.text) != a['text_hash']:
            raise ValueError(f"anchor text changed: {a['anchor_id']}")
        relation = op.get('insert_relation', 'AFTER')
        cur = cursors.get(a['anchor_id'], anchor)
        if op['page_break_before'] and op['orientation'] != 'LANDSCAPE':
            cur = add_page_break_after(cur)
        if op['orientation'] == 'LANDSCAPE':
            start_marker = insert_after(cur)
            set_section_end(start_marker, False)
            cur = start_marker
        if op.get('title'):
            cur = insert_after(cur, op['title'])
            try:
                cur.style = doc.styles['Heading 2']
            except Exception:
                pass
        if op['op_type'] == 'IMAGE':
            asset = get_asset(db_path, project_id, op['asset_id'])
            pimg = insert_after(cur) if relation == 'AFTER' else insert_before(cur)
            pimg.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pimg.add_run().add_picture(asset['file_path'], width=Cm(float(op.get('width_cm') or 16.0)))
            cur = pimg
            cap = op.get('caption') or f"{asset['asset_id']}｜{asset['kind']}｜{asset['qc_status']}"
            cur = insert_after(cur, cap)
            cur.alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            with connect(db_path) as con:
                vp = con.execute("SELECT content FROM video_prompts WHERE project_id=? AND scene_id=? AND status='FINAL' ORDER BY version DESC LIMIT 1", (project_id, op['scene_id'])).fetchone()
            if not vp:
                raise ValueError(f"FINAL video prompt missing for {op['scene_id']}")
            cur = insert_after(cur, vp['content']) if relation == 'AFTER' else insert_before(cur, vp['content'])
        if op['orientation'] == 'LANDSCAPE':
            end_marker = insert_after(cur)
            set_section_end(end_marker, True)
            cur = end_marker
        if op['page_break_after'] and op['orientation'] != 'LANDSCAPE':
            cur = add_page_break_after(cur)
        cursors[a['anchor_id']] = cur
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    with connect(db_path) as con:
        con.execute('UPDATE projects SET output_docx_path=?,release_status="PENDING",visual_page_qc_status="PENDING" WHERE project_id=?', (str(out), project_id))
        con.execute("DELETE FROM audits WHERE project_id=? AND audit_type IN ('PPI_MEDIA','RENDER_STRUCTURAL','VISUAL_PAGE_QC')", (project_id,))
    return out


def audit_docx_media(db_path: Path, project_id: str, docx_path: Path, target_ppi: float = 300, minimum_ppi: float = 240) -> dict[str, Any]:
    with connect(db_path) as con:
        manifest = {r['sha256']: dict(r) for r in con.execute('SELECT * FROM assets WHERE project_id=?', (project_id,))}
    rows = []
    errors = []
    with zipfile.ZipFile(docx_path) as z:
        relroot = ET.fromstring(z.read('word/_rels/document.xml.rels'))
        rels = {r.attrib.get('Id'): r.attrib for r in relroot}
        root = ET.fromstring(z.read('word/document.xml'))
        parents = root.findall('.//wp:inline', NS) + root.findall('.//wp:anchor', NS)
        for i, parent in enumerate(parents):
            blip = parent.find('.//a:blip', NS)
            if blip is None:
                continue
            rid = blip.attrib.get('{%s}embed' % NS['r'])
            link_rid = blip.attrib.get('{%s}link' % NS['r'])
            if link_rid:
                errors.append(f'external linked image {link_rid}')
                continue
            rel = rels.get(rid)
            if not rel:
                errors.append(f'missing rel {rid}')
                continue
            if rel.get('TargetMode') == 'External':
                errors.append(f'external rel {rid}')
                continue
            target = rel.get('Target', '').replace('../', '')
            member = 'word/' + target if not target.startswith('word/') else target
            try:
                raw = z.read(member)
            except KeyError:
                errors.append(f'missing media member: {member}')
                continue
            sha = hashlib.sha256(raw).hexdigest()
            asset = manifest.get(sha)
            ext = parent.find('wp:extent', NS)
            cx = int(ext.attrib.get('cx', '0')) if ext is not None else 0
            cy = int(ext.attrib.get('cy', '0')) if ext is not None else 0
            with Image.open(BytesIO(raw)) as im:
                px = im.size
            wi = cx / EMU_PER_IN if cx else 0
            hi = cy / EMU_PER_IN if cy else 0
            ppi = min(px[0] / wi if wi else 0, px[1] / hi if hi else 0)
            image_ratio = px[0] / px[1] if px[1] else 0
            display_ratio = wi / hi if hi else 0
            distortion = abs(display_ratio / image_ratio - 1) * 100 if image_ratio else 100
            if asset and asset['kind'] == 'DOCUMENT':
                status = 'BASELINE-PRESERVED' if ppi >= 150 and distortion <= 1 else 'BASELINE-WARN'
            else:
                status = 'PASS' if asset and asset['allowed_for_embed'] and ppi >= target_ppi and distortion <= 1 else ('DISCLOSED-MINIMUM' if asset and asset['allowed_for_embed'] and ppi >= minimum_ppi and distortion <= 1 else 'FAIL')
            if not asset:
                errors.append(f'unregistered embedded media: {member}')
            elif not asset['allowed_for_embed']:
                errors.append(f"not allowed_for_embed: {asset['asset_id']}")
            rows.append({'index': i, 'member': member, 'sha256': sha, 'asset_id': asset['asset_id'] if asset else None,
                         'kind': asset['kind'] if asset else None, 'scene_id': asset['scene_id'] if asset else None,
                         'pixels': list(px), 'display_inches': [round(wi, 3), round(hi, 3)],
                         'effective_ppi': round(ppi, 2), 'distortion_pct': round(distortion, 3), 'status': status})
        settings = z.read('word/settings.xml').decode('utf-8', 'ignore')
        no_compress = 'doNotCompressPictures' in settings
        m = re.search(r'defaultImageDpi[^>]+(?:val|w14:val)="(\d+)"', settings)
        default_dpi = int(m.group(1)) if m else None
        dpi_ok = default_dpi is not None and default_dpi >= 300
        external = [r for r in rels.values() if r.get('TargetMode') == 'External']
    required_assets = []
    with connect(db_path) as con:
        for r in con.execute("SELECT asset_id,kind FROM assets WHERE project_id=? AND allowed_for_embed=1 AND kind IN ('FORMAL_IMAGE','STORYBOARD_BOARD')", (project_id,)):
            required_assets.append(r['asset_id'])
    embedded_ids = [r['asset_id'] for r in rows if r['asset_id']]
    missing = sorted((Counter(required_assets) - Counter(embedded_ids)).elements())
    duplicates_required = sorted([aid for aid, count in Counter(embedded_ids).items() if aid in set(required_assets) and count > 1])
    if missing:
        errors.append(f'missing required embeds: {missing}')
    if duplicates_required:
        errors.append(f'duplicate required embeds: {duplicates_required}')
    if not no_compress:
        errors.append('doNotCompressPictures missing')
    if not dpi_ok:
        errors.append(f'defaultImageDpi invalid: {default_dpi}')
    report = {'images': rows, 'errors': errors, 'external_relationships': external,
              'doNotCompressPictures': no_compress, 'defaultImageDpi': default_dpi,
              'required_embeds': required_assets, 'embedded_asset_ids': embedded_ids,
              'missing_required': missing, 'duplicate_required_embeds': duplicates_required,
              'baseline_warnings': [r for r in rows if r['status'] == 'BASELINE-WARN'],
              'pass': not errors and not external and all(r['status'] in {'PASS','BASELINE-PRESERVED','BASELINE-WARN'} for r in rows)}
    return report
